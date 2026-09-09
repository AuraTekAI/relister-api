import io
import json
from datetime import datetime, timedelta
from itertools import count

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from rest_framework.test import APIClient

from .export_utils import days_to_sell_percentiles, quarter_date_range
from .models import Vehicle, VehicleListing
from .price_estimation import estimate_price, _parse_price, _mileage_matches
from .vehicle_export import _days_to_sell, _lifecycle_status

User = get_user_model()

# VehicleListing is unique per (user, list_id, seller_profile_id), so every
# fixture listing needs its own list_id.
_list_ids = count(1)


def _sold_listing(user, price, make="Toyota", model="Corolla", variant="Ascent Sport",
                  transmission="Automatic", year="2020", mileage=42000):
    vehicle = Vehicle.objects.create(
        make=make, model=model, variant=variant, transmission=transmission, year=year, mileage=mileage,
    )
    return VehicleListing.objects.create(
        user=user, vehicle=vehicle, price=price, status="sold",
        list_id=f"test-{next(_list_ids)}",
    )


class PriceEstimationUtilsTests(TestCase):
    def test_parse_price_handles_currency_formatting(self):
        self.assertEqual(_parse_price("$18,500"), 18500.0)
        self.assertEqual(_parse_price("18500"), 18500.0)

    def test_parse_price_rejects_unparsable_values(self):
        self.assertIsNone(_parse_price("POA"))
        self.assertIsNone(_parse_price(None))
        self.assertIsNone(_parse_price(""))
        self.assertIsNone(_parse_price("$0"))

    def test_mileage_window_is_500km_at_the_strict_tier(self):
        self.assertTrue(_mileage_matches(4400, 4000, 500))    # inside the window
        self.assertTrue(_mileage_matches(4500, 4000, 500))    # exactly on the edge
        self.assertFalse(_mileage_matches(4600, 4000, 500))   # just outside
        self.assertFalse(_mileage_matches(None, 4000, 500))   # unknown odometer never matches
        self.assertTrue(_mileage_matches(999_999, 4000, None))  # loosest tier ignores mileage


class PriceEstimationAlgorithmTests(TestCase):
    def setUp(self):
        self.dealer = User.objects.create_user(email="pricedealer@example.com", password="pw")

    def _criteria(self, **overrides):
        base = {
            "make": "Toyota", "model": "Corolla", "variant": "Ascent Sport",
            "transmission": "Automatic", "year": 2020, "mileage": 42000,
        }
        base.update(overrides)
        return base

    def test_exact_tier_used_when_enough_matches(self):
        for price in [15000, 16000, 17000, 18000, 19000]:
            _sold_listing(self.dealer, str(price))

        result = estimate_price(self._criteria())

        self.assertFalse(result["insufficient_data"])
        self.assertEqual(result["tier"], "exact")
        self.assertFalse(result["widened"])
        self.assertEqual(result["sample_size"], 5)
        self.assertEqual(result["estimated_price"], 17000.0)
        self.assertEqual(result["price_range"], {"low": 15500.0, "high": 18500.0})

    def test_widens_mileage_window_when_500km_is_too_tight(self):
        # Two sold within ±500 km of the target — not enough on its own.
        _sold_listing(self.dealer, "17000", mileage=42000)
        _sold_listing(self.dealer, "17500", mileage=42300)
        # Three more of the same car further out; only counted once the
        # mileage window widens to ±2,000 km.
        _sold_listing(self.dealer, "16000", mileage=43500)
        _sold_listing(self.dealer, "18000", mileage=43800)
        _sold_listing(self.dealer, "19000", mileage=40500)

        result = estimate_price(self._criteria(mileage=42000))

        self.assertFalse(result["insufficient_data"])
        self.assertEqual(result["tier"], "widen_mileage")
        self.assertTrue(result["widened"])
        self.assertEqual(result["sample_size"], 5)

    def test_widens_when_variant_does_not_match_enough_listings(self):
        # Only 2 sold with the exact variant — not enough on its own.
        _sold_listing(self.dealer, "17000", variant="Ascent Sport")
        _sold_listing(self.dealer, "17500", variant="Ascent Sport")
        # 3 more of the same car but a different (or unlabeled) variant —
        # only enough once variant is dropped from the match.
        _sold_listing(self.dealer, "16000", variant="SX")
        _sold_listing(self.dealer, "18000", variant="SX")
        _sold_listing(self.dealer, "19000", variant=None)

        result = estimate_price(self._criteria())

        self.assertFalse(result["insufficient_data"])
        self.assertEqual(result["tier"], "drop_variant")
        self.assertTrue(result["widened"])
        self.assertEqual(result["sample_size"], 5)

    def test_only_sold_listings_are_used(self):
        for price, status in [
            ("15000", "sold"),
            ("16000", "sold"),
            ("17000", "sold"),
            ("18000", "sold"),
            ("999999", "completed"),   # still listed — must be ignored
            ("1", "withdrawn"),        # withdrawn — must be ignored
            ("2", "pending"),          # pending — must be ignored
        ]:
            vehicle = Vehicle.objects.create(
                make="Toyota", model="Corolla", variant="Ascent Sport",
                transmission="Automatic", year="2020", mileage=42000,
            )
            VehicleListing.objects.create(
                user=self.dealer, vehicle=vehicle, price=price, status=status,
                list_id=f"test-{next(_list_ids)}",
            )

        # Only 4 genuinely sold — one short of the minimum, so even after
        # full widening this must report insufficient data, proving the
        # non-sold rows (which would otherwise push the count to 7) were
        # correctly excluded throughout every tier.
        result = estimate_price(self._criteria())
        self.assertTrue(result["insufficient_data"])
        self.assertEqual(result["sample_size"], 4)

    def test_insufficient_data_when_fewer_than_minimum_even_at_loosest_tier(self):
        _sold_listing(self.dealer, "15000")
        _sold_listing(self.dealer, "16000")

        result = estimate_price(self._criteria())

        self.assertTrue(result["insufficient_data"])
        self.assertIsNone(result["estimated_price"])
        self.assertIsNone(result["price_range"])
        self.assertEqual(result["sample_size"], 2)
        self.assertIn("2 matching sold vehicle", result["message"])

    def test_unparsable_prices_are_excluded_from_the_sample(self):
        for price in ["15000", "16000", "17000", "18000", "19000", "POA"]:
            _sold_listing(self.dealer, price)

        result = estimate_price(self._criteria())

        self.assertEqual(result["sample_size"], 5)

    def test_make_and_model_are_never_relaxed(self):
        for price in ["15000", "16000", "17000", "18000", "19000"]:
            _sold_listing(self.dealer, price, make="Honda", model="Civic")

        # Asking for a Toyota Corolla must never fall back to Honda Civic
        # sales, no matter how loose the other criteria get.
        result = estimate_price(self._criteria(make="Toyota", model="Corolla"))
        self.assertTrue(result["insufficient_data"])


class PriceEstimationEndpointTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(email="priceadmin@example.com", password="pw")
        self.admin.is_staff = True
        self.admin.save()
        self.dealer = User.objects.create_user(email="pricedealer2@example.com", password="pw")
        self.client = APIClient()
        self.url = reverse("estimate_vehicle_price")

        for price in [15000, 16000, 17000, 18000, 19000]:
            _sold_listing(self.dealer, str(price))

    def _post(self, payload, user=None):
        self.client.force_authenticate(user=user or self.admin)
        return self.client.post(self.url, data=json.dumps(payload), content_type="application/json")

    def test_returns_estimate_for_admin(self):
        response = self._post({
            "make": "Toyota", "model": "Corolla", "variant": "Ascent Sport",
            "transmission": "Automatic", "year": 2020, "mileage": 42000,
        })
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertFalse(body["insufficient_data"])
        self.assertEqual(body["estimated_price"], 17000.0)

    def test_rejects_non_admin(self):
        response = self._post(
            {"make": "Toyota", "model": "Corolla", "year": 2020, "mileage": 42000},
            user=self.dealer,
        )
        self.assertEqual(response.status_code, 403)

    def test_requires_make_and_model(self):
        response = self._post({"year": 2020, "mileage": 42000})
        self.assertEqual(response.status_code, 400)

    def test_requires_valid_year_and_mileage(self):
        response = self._post({"make": "Toyota", "model": "Corolla", "year": "not-a-year", "mileage": 42000})
        self.assertEqual(response.status_code, 400)

        response = self._post({"make": "Toyota", "model": "Corolla", "year": 2020, "mileage": -5})
        self.assertEqual(response.status_code, 400)


class VehicleExportTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(email="exportadmin@example.com", password="pw")
        self.admin.is_staff = True
        self.admin.save()
        self.dealer = User.objects.create_user(email="exportdealer@example.com", password="pw")
        self.dealer.dealership_name = "Best Cars Perth"
        self.dealer.save()

        self.client = APIClient()
        self.client.force_authenticate(user=self.admin)
        self.url = reverse("export_vehicle_data")

    def _make_sold_listing(self, days_to_sell, created_at=None):
        # days_to_sell is derived on this branch, so the fixture sets the two
        # timestamps it's computed from rather than the column directly.
        vehicle = Vehicle.objects.create(make="Toyota", model="Corolla", year="2020")
        sold_at = timezone.now()
        listing = VehicleListing.objects.create(
            user=self.dealer, vehicle=vehicle, price="15000", status="sold",
            listed_on=sold_at - timedelta(days=days_to_sell), sold_at=sold_at,
            list_id=f"export-{next(_list_ids)}",
        )
        if created_at:
            VehicleListing.objects.filter(pk=listing.pk).update(created_at=created_at)
        return listing

    def test_non_admin_is_forbidden(self):
        self.client.force_authenticate(user=self.dealer)
        response = self.client.get(self.url, {"export_format": "json"})
        self.assertEqual(response.status_code, 403)

    def test_json_export_includes_dealer_join_and_stats(self):
        self._make_sold_listing(10)
        self._make_sold_listing(20)
        self._make_sold_listing(30)

        response = self.client.get(self.url, {"export_format": "json"})
        self.assertEqual(response.status_code, 200)
        body = response.json()

        self.assertEqual(body["count"], 3)
        self.assertEqual(body["sold_count"], 3)
        self.assertEqual(body["stats"]["median_days_to_sell"], 20)

        row = body["results"][0]
        self.assertEqual(row["dealership_name"], "Best Cars Perth")
        self.assertEqual(row["dealer_email"], "exportdealer@example.com")
        self.assertEqual(row["make"], "Toyota")
        # The aggregate stat is repeated on every row.
        self.assertEqual(row["median_days_to_sell"], 20)

    def test_derived_lifecycle_columns_are_populated(self):
        listing = self._make_sold_listing(12)

        response = self.client.get(self.url, {"export_format": "json"})
        row = response.json()["results"][0]

        self.assertEqual(row["listing_id"], listing.id)
        self.assertEqual(row["lifecycle_status"], "sold")
        self.assertEqual(row["days_to_sell"], 12)
        self.assertIsNotNone(row["first_listed_at"])
        self.assertIsNotNone(row["delisted_at"])

    def test_non_sold_listings_excluded_from_stats(self):
        self._make_sold_listing(10)
        active_vehicle = Vehicle.objects.create(make="Honda", model="Civic", year="2021")
        VehicleListing.objects.create(
            user=self.dealer, vehicle=active_vehicle, price="20000", status="completed",
            list_id=f"export-{next(_list_ids)}",
        )

        response = self.client.get(self.url, {"export_format": "json"})
        body = response.json()
        self.assertEqual(body["count"], 2)       # both rows exported
        self.assertEqual(body["sold_count"], 1)  # only the sold one counted for stats

    def test_custom_date_range_filters_rows(self):
        in_range = timezone.make_aware(datetime(2025, 8, 15))
        out_of_range = timezone.make_aware(datetime(2025, 1, 1))
        self._make_sold_listing(5, created_at=in_range)
        self._make_sold_listing(5, created_at=out_of_range)

        response = self.client.get(self.url, {
            "export_format": "json", "start_date": "2025-07-01", "end_date": "2025-09-30",
        })
        body = response.json()
        self.assertEqual(body["count"], 1)

    def test_quarter_preset_matches_custom_range(self):
        in_q1 = timezone.make_aware(datetime(2025, 8, 1))
        self._make_sold_listing(5, created_at=in_q1)

        response = self.client.get(self.url, {
            "export_format": "json", "quarter": "Q1", "year": "2025", "year_type": "financial",
        })
        body = response.json()
        self.assertEqual(body["count"], 1)
        self.assertEqual(body["start_date"], "2025-07-01")
        self.assertEqual(body["end_date"], "2025-09-30")

    def test_xlsx_format_returns_a_real_spreadsheet(self):
        self._make_sold_listing(15)
        response = self.client.get(self.url, {"export_format": "xlsx"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment; filename=", response["Content-Disposition"])

        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual(sheet.cell(row=1, column=1).value, "Listing ID")
        self.assertEqual(sheet.max_row, 2)  # header + 1 data row

    def test_csv_format_returns_correct_headers(self):
        self._make_sold_listing(15)
        response = self.client.get(self.url, {"export_format": "csv"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/csv")
        first_line = response.content.decode().splitlines()[0]
        self.assertIn("Listing ID", first_line)
        self.assertIn("Dealership Name", first_line)

    def test_bad_format_is_rejected(self):
        response = self.client.get(self.url, {"export_format": "pdf"})
        self.assertEqual(response.status_code, 400)


class DerivedLifecycleColumnTests(TestCase):
    """The four columns this branch computes instead of storing."""

    def setUp(self):
        self.dealer = User.objects.create_user(email="derived@example.com", password="pw")

    def _listing(self, status, listed_on=None, sold_at=None):
        vehicle = Vehicle.objects.create(make="Toyota", model="Corolla", year="2020")
        return VehicleListing.objects.create(
            user=self.dealer, vehicle=vehicle, price="15000", status=status,
            listed_on=listed_on, sold_at=sold_at, list_id=f"derived-{next(_list_ids)}",
        )

    def test_lifecycle_status_maps_from_status(self):
        self.assertEqual(_lifecycle_status(self._listing("sold")), "sold")
        self.assertEqual(_lifecycle_status(self._listing("deleted")), "withdrawn")
        self.assertEqual(_lifecycle_status(self._listing("failed_deletion")), "withdrawn")
        self.assertEqual(_lifecycle_status(self._listing("completed")), "active")
        self.assertEqual(_lifecycle_status(self._listing("pending")), "active")

    def test_days_to_sell_counts_whole_days_between_listing_and_sale(self):
        sold_at = timezone.now()
        listing = self._listing("sold", listed_on=sold_at - timedelta(days=7), sold_at=sold_at)
        self.assertEqual(_days_to_sell(listing), 7)

    def test_days_to_sell_is_none_when_not_sold(self):
        sold_at = timezone.now()
        listing = self._listing("completed", listed_on=sold_at - timedelta(days=7), sold_at=sold_at)
        self.assertIsNone(_days_to_sell(listing))

    def test_days_to_sell_is_none_when_either_timestamp_is_missing(self):
        self.assertIsNone(_days_to_sell(self._listing("sold", sold_at=timezone.now())))
        self.assertIsNone(_days_to_sell(self._listing("sold", listed_on=timezone.now())))

    def test_days_to_sell_rejects_a_sale_before_the_listing(self):
        # An inconsistent pair must not report a negative duration, and must
        # not be silently reported as 0 days either.
        listed_on = timezone.now()
        listing = self._listing("sold", listed_on=listed_on, sold_at=listed_on - timedelta(days=3))
        self.assertIsNone(_days_to_sell(listing))


class ExportUtilsTests(TestCase):
    def test_australian_fy_q1_is_july_to_september(self):
        start, end = quarter_date_range("Q1", 2025, "financial")
        self.assertEqual(start.isoformat(), "2025-07-01")
        self.assertEqual(end.isoformat(), "2025-09-30")

    def test_australian_fy_q3_rolls_into_next_calendar_year(self):
        start, end = quarter_date_range("Q3", 2025, "financial")
        self.assertEqual(start.isoformat(), "2026-01-01")
        self.assertEqual(end.isoformat(), "2026-03-31")

    def test_calendar_year_q1_is_january_to_march(self):
        start, end = quarter_date_range("Q1", 2025, "calendar")
        self.assertEqual(start.isoformat(), "2025-01-01")
        self.assertEqual(end.isoformat(), "2025-03-31")

    def test_percentiles_empty_input_returns_none_not_zero(self):
        stats = days_to_sell_percentiles([])
        self.assertIsNone(stats["median_days_to_sell"])
